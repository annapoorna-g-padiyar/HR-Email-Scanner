import os
import re
import json

import cv2
import numpy as np
import pdfplumber
import pytesseract

from PIL import Image, ImageOps
from docx import Document


# ============================================================
# PROJECT PATHS
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

DATASET_DIR = os.path.join(
    BASE_DIR,
    "training_data"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "extracted_text"
)


RESUME_DIR = os.path.join(
    DATASET_DIR,
    "resumes"
)

NON_RESUME_DIR = os.path.join(
    DATASET_DIR,
    "non_resumes"
)


OUTPUT_RESUME_DIR = os.path.join(
    OUTPUT_DIR,
    "resume"
)

OUTPUT_NON_RESUME_DIR = os.path.join(
    OUTPUT_DIR,
    "non_resume"
)


os.makedirs(
    OUTPUT_RESUME_DIR,
    exist_ok=True
)

os.makedirs(
    OUTPUT_NON_RESUME_DIR,
    exist_ok=True
)


# ============================================================
# TESSERACT LOCATION
# ============================================================

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp"
}

DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt"
}


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = str(text)

    text = text.replace(
        "\x00",
        " "
    )

    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n+",
        "\n",
        text
    )

    return text.strip()


# ============================================================
# CHECK TEXT QUALITY
# ============================================================

def is_text_useful(text):

    if not text:
        return False

    text = text.strip()

    if len(text) < 100:
        return False

    words = re.findall(
        r"\b\w+\b",
        text
    )

    if len(words) < 20:
        return False

    return True


# ============================================================
# IMAGE PREPROCESSING
# ============================================================

def preprocess_image(image):

    image_np = np.array(image)

    # RGBA → RGB
    if len(image_np.shape) == 3:

        if image_np.shape[2] == 4:

            image_np = cv2.cvtColor(
                image_np,
                cv2.COLOR_RGBA2RGB
            )

        gray = cv2.cvtColor(
            image_np,
            cv2.COLOR_RGB2GRAY
        )

    else:

        gray = image_np


    # --------------------------------------------------------
    # UPSCALE SMALL IMAGES
    # --------------------------------------------------------

    height, width = gray.shape

    if width < 1800:

        scale = 1800 / width

        gray = cv2.resize(
            gray,
            None,
            fx=scale,
            fy=scale,
            interpolation=cv2.INTER_CUBIC
        )


    # --------------------------------------------------------
    # DENOISE
    # --------------------------------------------------------

    gray = cv2.GaussianBlur(
        gray,
        (3, 3),
        0
    )


    # --------------------------------------------------------
    # ADAPTIVE THRESHOLD
    # --------------------------------------------------------

    processed = cv2.adaptiveThreshold(
        gray,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        31,
        15
    )

    return processed


# ============================================================
# OCR
# ============================================================

def perform_ocr(image):

    try:

        processed = preprocess_image(
            image
        )

        text = pytesseract.image_to_string(
            processed,
            config="--oem 3 --psm 6"
        )

        return clean_text(
            text
        )

    except Exception as e:

        print(
            "OCR ERROR:",
            e
        )

        return ""


# ============================================================
# IMAGE → OCR
# ============================================================

def extract_image_text(file_path):

    try:

        image = Image.open(
            file_path
        )

        image = ImageOps.exif_transpose(
            image
        )

        image = image.convert(
            "RGB"
        )

        text = perform_ocr(
            image
        )

        return text

    except Exception as e:

        print(
            "IMAGE ERROR:",
            e
        )

        return ""


# ============================================================
# PDF NORMAL TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_path):

    text_parts = []

    try:

        with pdfplumber.open(
            file_path
        ) as pdf:

            for page in pdf.pages:

                text = page.extract_text()

                if text:

                    text_parts.append(
                        text
                    )

    except Exception as e:

        print(
            "PDF TEXT ERROR:",
            e
        )

        return ""


    return clean_text(
        "\n".join(
            text_parts
        )
    )


# ============================================================
# SCANNED PDF → OCR
# ============================================================

def ocr_pdf(file_path):

    text_parts = []

    try:

        with pdfplumber.open(
            file_path
        ) as pdf:

            for page_number, page in enumerate(
                pdf.pages,
                start=1
            ):

                print(
                    f"       OCR page {page_number}"
                )

                image = page.to_image(
                    resolution=250
                ).original

                image = image.convert(
                    "RGB"
                )

                page_text = perform_ocr(
                    image
                )

                if page_text:

                    text_parts.append(
                        page_text
                    )

    except Exception as e:

        print(
            "PDF OCR ERROR:",
            e
        )

        return ""


    return clean_text(
        "\n".join(
            text_parts
        )
    )


# ============================================================
# PROCESS PDF
# ============================================================

def process_pdf(file_path):

    print(
        "    Trying normal PDF text extraction..."
    )

    text = extract_pdf_text(
        file_path
    )


    # Selectable PDF
    if is_text_useful(text):

        print(
            "    Method: Normal PDF text extraction"
        )

        return text, "pdf_text"


    # Scanned PDF
    print(
        "    Little/no text found."
    )

    print(
        "    Switching to OCR..."
    )

    text = ocr_pdf(
        file_path
    )

    return text, "pdf_ocr"


# ============================================================
# DOCX
# ============================================================

def extract_docx_text(file_path):

    text_parts = []

    try:

        document = Document(
            file_path
        )


        # Paragraphs
        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text_parts.append(
                    paragraph.text
                )


        # Tables
        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():

                        text_parts.append(
                            cell.text
                        )


    except Exception as e:

        print(
            "DOCX ERROR:",
            e
        )

        return ""


    return clean_text(
        "\n".join(
            text_parts
        )
    )


# ============================================================
# TXT
# ============================================================

def extract_txt_text(file_path):

    try:

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            text = file.read()

        return clean_text(
            text
        )

    except Exception as e:

        print(
            "TXT ERROR:",
            e
        )

        return ""


# ============================================================
# MASTER EXTRACTION FUNCTION
# ============================================================

def extract_file(file_path):

    extension = os.path.splitext(
        file_path
    )[1].lower()


    # PDF
    if extension == ".pdf":

        return process_pdf(
            file_path
        )


    # DOCX
    elif extension == ".docx":

        text = extract_docx_text(
            file_path
        )

        return text, "docx"


    # TXT
    elif extension == ".txt":

        text = extract_txt_text(
            file_path
        )

        return text, "txt"


    # Images
    elif extension in IMAGE_EXTENSIONS:

        text = extract_image_text(
            file_path
        )

        return text, "ocr_image"


    return "", "unsupported"


# ============================================================
# PROCESS ONE CLASS
# ============================================================

def process_class(
    class_name,
    input_dir,
    output_dir
):

    print("\n")
    print("=" * 70)

    print(
        f"PROCESSING {class_name.upper()}"
    )

    print("=" * 70)


    if not os.path.exists(
        input_dir
    ):

        print(
            "ERROR: Folder not found:"
        )

        print(
            input_dir
        )

        return


    files = os.listdir(
        input_dir
    )


    total = 0
    successful = 0
    failed = 0


    for filename in files:

        file_path = os.path.join(
            input_dir,
            filename
        )


        if not os.path.isfile(
            file_path
        ):

            continue


        extension = os.path.splitext(
            filename
        )[1].lower()


        if (
            extension not in IMAGE_EXTENSIONS
            and extension not in DOCUMENT_EXTENSIONS
        ):

            continue


        total += 1


        print("\n")
        print("-" * 70)

        print(
            f"FILE: {filename}"
        )

        print("-" * 70)


        text, method = extract_file(
            file_path
        )


        text = clean_text(
            text
        )


        # ----------------------------------------------------
        # Use filename safely
        # ----------------------------------------------------

        base_name = os.path.splitext(
            filename
        )[0]

        # Remove unsafe filename characters
        safe_name = re.sub(
            r'[<>:"/\\|?*]',
            "_",
            base_name
        )


        output_file = os.path.join(
            output_dir,
            safe_name + ".txt"
        )


        metadata_file = os.path.join(
            output_dir,
            safe_name + "_metadata.json"
        )


        # ----------------------------------------------------
        # Save extracted text
        # ----------------------------------------------------

        with open(
            output_file,
            "w",
            encoding="utf-8"
        ) as file:

            file.write(
                text
            )


        # ----------------------------------------------------
        # Metadata
        # ----------------------------------------------------

        words = re.findall(
            r"\b\w+\b",
            text
        )


        metadata = {

            "original_file": filename,

            "class": class_name,

            "extraction_method": method,

            "characters": len(text),

            "words": len(words),

            "usable_text": is_text_useful(text)

        }


        with open(
            metadata_file,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                metadata,
                file,
                indent=4
            )


        # ----------------------------------------------------
        # Result
        # ----------------------------------------------------

        print(
            f"    Method: {method}"
        )

        print(
            f"    Characters: {len(text)}"
        )

        print(
            f"    Words: {len(words)}"
        )


        if is_text_useful(text):

            print(
                "    Status: SUCCESS"
            )

            successful += 1

        else:

            print(
                "    Status: WARNING - "
                "insufficient text"
            )

            failed += 1


    print("\n")
    print(
        f"{class_name.upper()} SUMMARY"
    )

    print(
        f"Total files : {total}"
    )

    print(
        f"Successful  : {successful}"
    )

    print(
        f"Poor/failed : {failed}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("\n")
    print("=" * 70)
    print("RESUME DATASET EXTRACTION + OCR")
    print("=" * 70)


    # --------------------------------------------------------
    # Check dataset
    # --------------------------------------------------------

    if not os.path.exists(
        DATASET_DIR
    ):

        print(
            "\nERROR: training data folder not found:"
        )

        print(
            DATASET_DIR
        )

        return


    # --------------------------------------------------------
    # Check Tesseract
    # --------------------------------------------------------

    if not os.path.exists(
        TESSERACT_PATH
    ):

        print(
            "\nERROR: Tesseract OCR not found."
        )

        print(
            "Expected:"
        )

        print(
            TESSERACT_PATH
        )

        return


    # --------------------------------------------------------
    # Process resumes
    # --------------------------------------------------------

    process_class(

        "resume",

        RESUME_DIR,

        OUTPUT_RESUME_DIR

    )


    # --------------------------------------------------------
    # Process non-resumes
    # --------------------------------------------------------

    process_class(

        "non_resume",

        NON_RESUME_DIR,

        OUTPUT_NON_RESUME_DIR

    )


    print("\n")
    print("=" * 70)
    print("EXTRACTION COMPLETED")
    print("=" * 70)

    print(
        "\nOutput:"
    )

    print(
        OUTPUT_DIR
    )


if __name__ == "__main__":

    main()