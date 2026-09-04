import os
import re
import shutil
import glob
import joblib

import cv2
import pytesseract
import pdfplumber

# ============================================================
# OPTIONAL DOCX SUPPORT
# ============================================================

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


# ============================================================
# TESSERACT
# ============================================================

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# ============================================================
# PROJECT PATH
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)


# ============================================================
# DIRECTORIES
# ============================================================

ATTACHMENT_DIR = os.path.join(
    BASE_DIR,
    "attachments"
)

CLASSIFIED_DIR = os.path.join(
    BASE_DIR,
    "classified"
)

RESUME_OUTPUT = os.path.join(
    CLASSIFIED_DIR,
    "resumes"
)

NON_RESUME_OUTPUT = os.path.join(
    CLASSIFIED_DIR,
    "non_resumes"
)

REVIEW_OUTPUT = os.path.join(
    CLASSIFIED_DIR,
    "review"
)


os.makedirs(
    RESUME_OUTPUT,
    exist_ok=True
)

os.makedirs(
    NON_RESUME_OUTPUT,
    exist_ok=True
)

os.makedirs(
    REVIEW_OUTPUT,
    exist_ok=True
)


# ============================================================
# MODEL
# ============================================================

MODEL_FILE = os.path.join(
    BASE_DIR,
    "models",
    "resume_classifier.pkl"
)


# ============================================================
# RESUME SECTIONS
# ============================================================

RESUME_SECTIONS = {

    "objective": [
        "objective",
        "career objective",
        "professional objective"
    ],

    "summary": [
        "professional summary",
        "profile summary",
        "career summary",
        "summary"
    ],

    "education": [
        "education",
        "educational qualification",
        "academic qualification",
        "academic background",
        "qualification"
    ],

    "experience": [
        "work experience",
        "professional experience",
        "employment history",
        "experience",
        "work history"
    ],

    "skills": [
        "technical skills",
        "skills",
        "technical expertise",
        "core skills",
        "key skills"
    ],

    "projects": [
        "projects",
        "academic projects",
        "personal projects",
        "project experience"
    ],

    "certifications": [
        "certifications",
        "certificates",
        "professional certifications"
    ],

    "internship": [
        "internship",
        "internships",
        "intern experience"
    ]
}


# ============================================================
# GENERAL RESUME INDICATORS
# ============================================================

GENERAL_RESUME_INDICATORS = [

    "resume",
    "curriculum vitae",
    "cv",

    "linkedin",
    "github",
    "portfolio",

    "python",
    "java",
    "javascript",
    "c++",
    "c#",
    "sql",
    "html",
    "css",

    "developer",
    "software engineer",
    "engineer",
    "programmer",

    "b.tech",
    "btech",
    "b.e",
    "m.tech",
    "mtech",
    "m.e",

    "bachelor",
    "master",
    "degree",

    "internship",
    "intern",

    "experience",
    "projects",
    "skills",

    "gpa",
    "cgpa"
]


# ============================================================
# NON-RESUME INDICATORS
# ============================================================

NON_RESUME_INDICATORS = [

    "event",
    "workshop",
    "workshops",
    "seminar",
    "conference",
    "webinar",

    "invitation",
    "invited",
    "you're invited",
    "you are invited",
    "rsvp",
    "please rsvp",
    "reserve your spot",
    "reserve your seat",
    "save your spot",

    "employee training",
    "training workshop",
    "training session",
    "training program",
    "development workshop",

    "team building",
    "team-building",

    "marks card",
    "mark card",
    "mark sheet",
    "marksheet",
    "grade card",
    "report card",
    "academic record",
    "transcript",

    "admit card",
    "hall ticket",
    "question paper",
    "exam schedule",
    "examination",

    "invoice",
    "receipt",
    "purchase order",
    "quotation",

    "brochure",
    "newsletter",
    "memorandum",
    "memo",
    "registration form",
    "application form"
]


# ============================================================
# STRONG NON-RESUME PHRASES
# ============================================================

STRONG_NON_RESUME_PHRASES = [

    "workshop",
    "training workshop",
    "employee training",
    "training session",
    "training program",
    "development workshop",

    "invitation",
    "you're invited",
    "you are invited",

    "rsvp",
    "please rsvp",

    "reserve your spot",
    "reserve your seat",
    "save your spot",

    "conference",
    "seminar",
    "webinar",

    "event registration",
    "event schedule",
    "event agenda",

    "meeting invitation",
    "meeting agenda",

    "marks card",
    "mark card",
    "mark sheet",
    "marksheet",

    "grade card",
    "report card",
    "academic record",
    "transcript",

    "admit card",
    "hall ticket",
    "question paper",

    "invoice",
    "receipt",

    "registration form",
    "application form",

    "brochure",
    "newsletter"
]


# ============================================================
# STRONG RESUME PHRASES
# ============================================================

STRONG_RESUME_PHRASES = [

    "resume",
    "curriculum vitae",

    "professional summary",
    "career objective",

    "work experience",
    "professional experience",
    "employment history",

    "technical skills",

    "educational qualification",
    "academic qualification",

    "project experience",

    "certifications",

    "internship",
    "internships"
]


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = text.lower()

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# WHOLE WORD MATCHING
# ============================================================

def contains_keyword(text, keyword):

    keyword = keyword.strip()

    if not keyword:
        return False

    pattern = (
        r"(?<!\w)"
        + re.escape(keyword)
        + r"(?!\w)"
    )

    return re.search(
        pattern,
        text,
        re.IGNORECASE
    ) is not None


# ============================================================
# TESSERACT CHECK
# ============================================================

print()
print("=" * 70)
print("HR RESUME / NON-RESUME CLASSIFIER")
print("=" * 70)

print()
print("Checking Tesseract...")

if not os.path.exists(TESSERACT_PATH):

    print("ERROR: Tesseract not found:")
    print(TESSERACT_PATH)

    raise SystemExit

try:

    version = pytesseract.get_tesseract_version()

    print(
        f"Tesseract detected: {version}"
    )

except Exception as e:

    print(
        "ERROR: Tesseract could not be used."
    )

    print(e)

    raise SystemExit


# ============================================================
# LOAD MODEL
# ============================================================

print()
print("Loading classifier...")

if not os.path.exists(MODEL_FILE):

    print("ERROR: Model not found:")
    print(MODEL_FILE)

    print()
    print("Train the model first.")

    raise SystemExit


try:

    model = joblib.load(
        MODEL_FILE
    )

    print("Existing classifier loaded successfully.")

except Exception as e:

    print("ERROR loading model:")
    print(e)

    raise SystemExit


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_path):

    text_parts = []

    try:

        with pdfplumber.open(
            file_path
        ) as pdf:

            print(
                f"    PDF pages: {len(pdf.pages)}"
            )

            for page in pdf.pages:

                page_text = page.extract_text()

                if page_text:

                    text_parts.append(
                        page_text
                    )

    except Exception as e:

        print(
            "    PDF extraction error:",
            e
        )

    return clean_text(
        "\n".join(text_parts)
    )


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(file_path):

    if not DOCX_AVAILABLE:

        print(
            "    python-docx is not installed."
        )

        return ""

    text_parts = []

    try:

        document = Document(
            file_path
        )

        for paragraph in document.paragraphs:

            if paragraph.text:

                text_parts.append(
                    paragraph.text
                )

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    if cell.text:

                        row_text.append(
                            cell.text
                        )

                if row_text:

                    text_parts.append(
                        " ".join(row_text)
                    )

    except Exception as e:

        print(
            "    DOCX extraction error:",
            e
        )

    return clean_text(
        "\n".join(text_parts)
    )


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_txt_text(file_path):

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1"
    ]

    for encoding in encodings:

        try:

            with open(
                file_path,
                "r",
                encoding=encoding,
                errors="ignore"
            ) as file:

                return clean_text(
                    file.read()
                )

        except Exception:

            continue

    return ""


# ============================================================
# IMAGE OCR
# ============================================================

def ocr_image(file_path):

    image = cv2.imread(
        file_path
    )

    if image is None:

        print(
            "    Could not read image."
        )

        return ""

    height, width = image.shape[:2]

    max_dimension = 3000

    scale = min(
        max_dimension / width,
        max_dimension / height
    )

    if scale > 1:

        image = cv2.resize(
            image,
            None,
            fx=scale,
            fy=scale,
            interpolation=cv2.INTER_CUBIC
        )

    gray = cv2.cvtColor(
        image,
        cv2.COLOR_BGR2GRAY
    )

    results = []

    # --------------------------------------------------------
    # OCR 1
    # --------------------------------------------------------

    try:

        text = pytesseract.image_to_string(
            gray,
            config="--psm 6"
        )

        if text:

            results.append(
                text
            )

    except Exception:
        pass

    # --------------------------------------------------------
    # OCR 2
    # --------------------------------------------------------

    try:

        text = pytesseract.image_to_string(
            gray,
            config="--psm 11"
        )

        if text:

            results.append(
                text
            )

    except Exception:
        pass

    # --------------------------------------------------------
    # OTSU
    # --------------------------------------------------------

    try:

        _, threshold = cv2.threshold(
            gray,
            0,
            255,
            cv2.THRESH_BINARY + cv2.THRESH_OTSU
        )

        text = pytesseract.image_to_string(
            threshold,
            config="--psm 11"
        )

        if text:

            results.append(
                text
            )

    except Exception:
        pass

    # --------------------------------------------------------
    # ADAPTIVE
    # --------------------------------------------------------

    try:

        adaptive = cv2.adaptiveThreshold(
            gray,
            255,
            cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
            cv2.THRESH_BINARY,
            31,
            15
        )

        text = pytesseract.image_to_string(
            adaptive,
            config="--psm 11"
        )

        if text:

            results.append(
                text
            )

    except Exception:
        pass

    # --------------------------------------------------------
    # BEST OCR RESULT
    # --------------------------------------------------------

    best_text = ""

    best_score = 0

    for result in results:

        cleaned = clean_text(
            result
        )

        words = cleaned.split()

        score = (
            len(cleaned)
            + len(words) * 3
        )

        if score > best_score:

            best_score = score

            best_text = cleaned

    return best_text


# ============================================================
# PDF OCR
# ============================================================

def ocr_pdf(file_path):

    text_parts = []

    try:

        import fitz

        document = fitz.open(
            file_path
        )

        print(
            f"    OCR processing {len(document)} page(s)..."
        )

        for page_number, page in enumerate(
            document,
            start=1
        ):

            print(
                f"       OCR page {page_number}..."
            )

            pix = page.get_pixmap(
                matrix=fitz.Matrix(
                    2,
                    2
                ),
                alpha=False
            )

            temp_dir = os.environ.get(
                "TEMP",
                "."
            )

            image_path = os.path.join(
                temp_dir,
                f"hr_ocr_{os.getpid()}_{page_number}.png"
            )

            pix.save(
                image_path
            )

            page_text = ocr_image(
                image_path
            )

            if page_text:

                text_parts.append(
                    page_text
                )

            try:

                os.remove(
                    image_path
                )

            except Exception:
                pass

        document.close()

    except Exception as e:

        print(
            "    PDF OCR error:",
            e
        )

    return clean_text(
        " ".join(text_parts)
    )


# ============================================================
# EXTRACT DOCUMENT
# ============================================================

def extract_document(file_path):

    extension = os.path.splitext(
        file_path
    )[1].lower()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        print(
            "    Reading PDF..."
        )

        text = extract_pdf_text(
            file_path
        )

        characters = len(text)

        words = len(
            text.split()
        )

        if characters >= 100 and words >= 20:

            return (
                text,
                "pdf_text"
            )

        print(
            "    Very little usable PDF text."
        )

        print(
            "    Using OCR..."
        )

        text = ocr_pdf(
            file_path
        )

        return (
            text,
            "pdf_ocr"
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == ".docx":

        print(
            "    Reading DOCX..."
        )

        return (
            extract_docx_text(
                file_path
            ),
            "docx_text"
        )

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    if extension == ".txt":

        print(
            "    Reading TXT..."
        )

        return (
            extract_txt_text(
                file_path
            ),
            "txt"
        )

    # --------------------------------------------------------
    # IMAGE
    # --------------------------------------------------------

    if extension in [
        ".jpg",
        ".jpeg",
        ".png",
        ".bmp",
        ".tif",
        ".tiff"
    ]:

        print(
            "    Reading image..."
        )

        return (
            ocr_image(
                file_path
            ),
            "ocr_image"
        )

    return (
        "",
        "unsupported"
    )


# ============================================================
# FIND RESUME SECTIONS
# ============================================================

def find_resume_sections(text):

    found = []

    for section, keywords in RESUME_SECTIONS.items():

        for keyword in keywords:

            if contains_keyword(
                text,
                keyword
            ):

                found.append(
                    section
                )

                break

    return list(
        set(found)
    )


# ============================================================
# GENERAL RESUME INDICATORS
# ============================================================

def count_general_indicators(text):

    matches = []

    for keyword in GENERAL_RESUME_INDICATORS:

        if contains_keyword(
            text,
            keyword
        ):

            matches.append(
                keyword
            )

    return list(
        set(matches)
    )


# ============================================================
# NON-RESUME INDICATORS
# ============================================================

def count_negative_indicators(text):

    matches = []

    for keyword in NON_RESUME_INDICATORS:

        if contains_keyword(
            text,
            keyword
        ):

            matches.append(
                keyword
            )

    return list(
        set(matches)
    )


# ============================================================
# STRONG NON-RESUME
# ============================================================

def find_strong_non_resume(text):

    matches = []

    for phrase in STRONG_NON_RESUME_PHRASES:

        if contains_keyword(
            text,
            phrase
        ):

            matches.append(
                phrase
            )

    return list(
        set(matches)
    )


# ============================================================
# STRONG RESUME
# ============================================================

def find_strong_resume(text):

    matches = []

    for phrase in STRONG_RESUME_PHRASES:

        if contains_keyword(
            text,
            phrase
        ):

            matches.append(
                phrase
            )

    return list(
        set(matches)
    )


# ============================================================
# TEXT QUALITY
# ============================================================

def get_text_quality(text):

    characters = len(text)

    words = len(
        text.split()
    )

    if characters < 150 or words < 25:

        return "POOR"

    if characters < 500 or words < 80:

        return "MEDIUM"

    return "GOOD"


# ============================================================
# UNIQUE DESTINATION
# ============================================================

def get_unique_destination(
    directory,
    filename
):

    base, extension = os.path.splitext(
        filename
    )

    destination = os.path.join(
        directory,
        filename
    )

    counter = 1

    while os.path.exists(
        destination
    ):

        destination = os.path.join(
            directory,
            f"{base}_{counter}{extension}"
        )

        counter += 1

    return destination


# ============================================================
# FINAL DECISION
# ============================================================

def make_final_decision(
    ml_prediction,
    confidence,
    text_quality,
    section_score,
    general_matches,
    negative_matches,
    text
):

    normalized = clean_text(
        text
    )

    strong_non_resume_found = (
        find_strong_non_resume(
            normalized
        )
    )

    strong_non_resume_count = len(
        strong_non_resume_found
    )

    strong_resume_found = (
        find_strong_resume(
            normalized
        )
    )

    strong_resume_count = len(
        strong_resume_found
    )

    characters = len(
        normalized
    )

    words = len(
        normalized.split()
    )

    # ========================================================
    # STRONG RESUME STRUCTURE
    # ========================================================

    if section_score >= 3:

        if strong_non_resume_count >= 2:

            return (
                "NON-RESUME",
                "Strong non-resume evidence"
            )

        return (
            "RESUME",
            "Strong resume section structure"
        )

    # ========================================================
    # MULTIPLE NON-RESUME INDICATORS
    # ========================================================

    if strong_non_resume_count >= 2:

        if section_score <= 1:

            return (
                "NON-RESUME",
                "Multiple non-resume indicators"
            )

    # ========================================================
    # STRONG NON-RESUME INDICATOR
    # ========================================================

    if strong_non_resume_count >= 1:

        if section_score == 0:

            return (
                "NON-RESUME",
                "Strong non-resume indicator"
            )

    # ========================================================
    # STRONG RESUME INDICATORS
    # ========================================================

    if strong_resume_count >= 3:

        if section_score >= 1:

            return (
                "RESUME",
                "Multiple resume indicators"
            )

    # ========================================================
    # RESUME STRUCTURE
    # ========================================================

    if section_score >= 2:

        if confidence >= 50:

            return (
                "RESUME",
                "Resume sections detected"
            )

    # ========================================================
    # HIGH CONFIDENCE ML
    # ========================================================

    if confidence >= 85:

        if ml_prediction == 1:

            if strong_non_resume_count == 0:

                return (
                    "RESUME",
                    "High-confidence ML prediction"
                )

        else:

            return (
                "NON-RESUME",
                "High-confidence ML prediction"
            )

    # ========================================================
    # MEDIUM CONFIDENCE
    # ========================================================

    if confidence >= 65:

        if ml_prediction == 1:

            if (
                section_score >= 1
                or general_matches >= 3
            ):

                return (
                    "RESUME",
                    "ML prediction supported by resume indicators"
                )

        else:

            if section_score == 0:

                return (
                    "NON-RESUME",
                    "ML prediction with no resume structure"
                )

    # ========================================================
    # LOW CONFIDENCE
    # ========================================================

    if confidence < 65:

        if ml_prediction == 1:

            if section_score >= 2:

                return (
                    "RESUME",
                    "Resume structure despite low ML confidence"
                )

        if ml_prediction == 0:

            if section_score == 0:

                return (
                    "NON-RESUME",
                    "No resume structure"
                )

    # ========================================================
    # SHORT DOCUMENT
    # ========================================================

    if characters < 150 or words < 25:

        if section_score == 0:

            if strong_resume_count == 0:

                if strong_non_resume_count >= 1:

                    return (
                        "NON-RESUME",
                        "Short document with non-resume indicator"
                    )

                if ml_prediction == 0:

                    return (
                        "NON-RESUME",
                        "Short document with no resume structure"
                    )

                return (
                    "NON-RESUME",
                    "Very short document with no resume structure"
                )

    # ========================================================
    # POOR OCR
    # ========================================================

    if text_quality == "POOR":

        if section_score >= 1:

            return (
                "RESUME",
                "Resume indicator in low-text OCR"
            )

        if strong_non_resume_count >= 1:

            return (
                "NON-RESUME",
                "Non-resume indicator in low-text OCR"
            )

    # ========================================================
    # NON-RESUME INDICATORS
    # ========================================================

    if section_score == 0:

        if negative_matches >= 1:

            return (
                "NON-RESUME",
                "Non-resume indicators detected"
            )

    # ========================================================
    # FALLBACK
    # ========================================================

    return (
        "REVIEW",
        "Evidence is insufficient or conflicting"
    )


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_EXTENSIONS = [

    ".pdf",

    ".docx",

    ".txt",

    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff"
]


# ============================================================
# FIND ATTACHMENTS
# ============================================================

files = []

for extension in SUPPORTED_EXTENSIONS:

    files.extend(
        glob.glob(
            os.path.join(
                ATTACHMENT_DIR,
                "*" + extension
            )
        )
    )


files = list(
    dict.fromkeys(
        files
    )
)


# ============================================================
# START
# ============================================================

print()
print("=" * 70)
print("CLASSIFYING ATTACHMENTS")
print("=" * 70)

print(
    f"\nTotal attachments found: {len(files)}"
)


# ============================================================
# COUNTERS
# ============================================================

resume_count = 0

non_resume_count = 0

review_count = 0

failure_count = 0


# ============================================================
# PROCESS
# ============================================================

for index, file_path in enumerate(
    files,
    start=1
):

    filename = os.path.basename(
        file_path
    )

    print()
    print("-" * 70)

    print(
        f"[{index}/{len(files)}] {filename}"
    )

    print("-" * 70)

    # ========================================================
    # EXTRACTION
    # ========================================================

    try:

        text, method = extract_document(
            file_path
        )

    except Exception as e:

        print(
            "    Extraction failed:",
            e
        )

        failure_count += 1

        destination = get_unique_destination(
            REVIEW_OUTPUT,
            filename
        )

        try:

            shutil.copy2(
                file_path,
                destination
            )

        except Exception:
            pass

        review_count += 1

        print(
            "    RESULT: REVIEW"
        )

        continue

    # ========================================================
    # TEXT INFORMATION
    # ========================================================

    characters = len(text)

    words = len(
        text.split()
    )

    print(
        f"    Characters : {characters}"
    )

    print(
        f"    Words      : {words}"
    )

    print(
        f"    Method     : {method}"
    )

    # ========================================================
    # NO TEXT
    # ========================================================

    if not text:

        destination = get_unique_destination(
            REVIEW_OUTPUT,
            filename
        )

        try:

            shutil.copy2(
                file_path,
                destination
            )

        except Exception:

            failure_count += 1

        review_count += 1

        print(
            "    RESULT     : REVIEW"
        )

        print(
            "    Reason     : No text extracted"
        )

        continue

    # ========================================================
    # TEXT QUALITY
    # ========================================================

    text_quality = get_text_quality(
        text
    )

    print(
        f"    Text quality: {text_quality}"
    )

    # ========================================================
    # RESUME SECTIONS
    # ========================================================

    sections = find_resume_sections(
        text
    )

    section_score = len(
        sections
    )

    print(
        f"    Resume sections: {section_score}"
    )

    # ========================================================
    # INDICATOR COUNTS
    # ========================================================

    general_matches_list = (
        count_general_indicators(
            text
        )
    )

    general_matches = len(
        general_matches_list
    )

    negative_matches_list = (
        count_negative_indicators(
            text
        )
    )

    negative_matches = len(
        negative_matches_list
    )

    # ========================================================
    # ML PREDICTION
    # ========================================================

    try:

        prediction = model.predict(
            [text]
        )[0]

        probabilities = model.predict_proba(
            [text]
        )[0]

        confidence = (
            max(probabilities)
            * 100
        )

        prediction_name = (
            "RESUME"
            if prediction == 1
            else "NON-RESUME"
        )

        print(
            f"    ML prediction : {prediction_name}"
        )

        print(
            f"    Confidence    : {confidence:.2f}%"
        )

    except Exception as e:

        print(
            "    ML prediction failed:",
            e
        )

        destination = get_unique_destination(
            REVIEW_OUTPUT,
            filename
        )

        try:

            shutil.copy2(
                file_path,
                destination
            )

        except Exception:
            pass

        failure_count += 1

        review_count += 1

        print(
            "    RESULT        : REVIEW"
        )

        continue

    # ========================================================
    # FINAL DECISION
    # ========================================================

    final_result, reason = (
        make_final_decision(
            prediction,
            confidence,
            text_quality,
            section_score,
            general_matches,
            negative_matches,
            text
        )
    )

    # ========================================================
    # OUTPUT DIRECTORY
    # ========================================================

    if final_result == "RESUME":

        destination_dir = (
            RESUME_OUTPUT
        )

        resume_count += 1

    elif final_result == "NON-RESUME":

        destination_dir = (
            NON_RESUME_OUTPUT
        )

        non_resume_count += 1

    else:

        destination_dir = (
            REVIEW_OUTPUT
        )

        review_count += 1

    # ========================================================
    # COPY
    # ========================================================

    destination = get_unique_destination(
        destination_dir,
        filename
    )

    try:

        shutil.copy2(
            file_path,
            destination
        )

        print()
        print(
            f"    RESULT : {final_result}"
        )

        print(
            f"    Reason : {reason}"
        )

        print(
            f"    Saved  : {destination}"
        )

    except Exception as e:

        print(
            "    Could not copy file:",
            e
        )

        failure_count += 1


# ============================================================
# FINAL SUMMARY
# ============================================================

print()
print()
print("=" * 70)
print("CLASSIFICATION COMPLETED")
print("=" * 70)

print(
    f"Total attachments : {len(files)}"
)

print(
    f"Resumes           : {resume_count}"
)

print(
    f"Non-resumes       : {non_resume_count}"
)

print(
    f"Review             : {review_count}"
)

print(
    f"Extraction/ML fail : {failure_count}"
)

print()
print(
    "Resume folder:"
)

print(
    RESUME_OUTPUT
)

print()
print(
    "Non-resume folder:"
)

print(
    NON_RESUME_OUTPUT
)

print()
print(
    "Review folder:"
)

print(
    REVIEW_OUTPUT
)

print()
print(
    "Original attachments were NOT modified."
)

print("=" * 70)