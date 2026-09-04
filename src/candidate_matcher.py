import os
import re
import json
from pathlib import Path

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(r"C:\Users\LENOVO\HR_Resume_AI")

RESUME_FOLDER = BASE_DIR / "classified" / "resumes"
RESULTS_FOLDER = BASE_DIR / "results"
RESULT_FILE = RESULTS_FOLDER / "candidate_ranking.json"

RESULTS_FOLDER.mkdir(parents=True, exist_ok=True)


# ============================================================
# OPTIONAL LIBRARIES
# ============================================================

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    from pdf2image import convert_from_path
except ImportError:
    convert_from_path = None


# ============================================================
# SUPPORTED FILES
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".png",
    ".jpg",
    ".jpeg"
}


# ============================================================
# SKILL DATABASE
# ============================================================

SKILLS = [

    # --------------------------------------------------------
    # PROGRAMMING LANGUAGES
    # --------------------------------------------------------

    "python",
    "java",
    "c++",
    "c#",
    "javascript",
    "typescript",
    "php",
    "ruby",
    "go",
    "kotlin",
    "swift",

    # --------------------------------------------------------
    # WEB DEVELOPMENT
    # --------------------------------------------------------

    "html",
    "css",
    "bootstrap",
    "react",
    "react.js",
    "angular",
    "vue",
    "vue.js",
    "node",
    "node.js",
    "express",
    "express.js",
    "django",
    "flask",
    "spring",
    "spring boot",

    # --------------------------------------------------------
    # DATABASES
    # --------------------------------------------------------

    "sql",
    "mysql",
    "postgresql",
    "postgres",
    "mongodb",
    "oracle",
    "sqlite",
    "redis",

    # --------------------------------------------------------
    # API
    # --------------------------------------------------------

    "rest api",
    "rest apis",
    "restful api",
    "restful apis",
    "web api",
    "web apis",
    "api",
    "apis",
    "graphql",

    # --------------------------------------------------------
    # VERSION CONTROL / DEVOPS
    # --------------------------------------------------------

    "git",
    "github",
    "gitlab",
    "docker",
    "kubernetes",
    "jenkins",
    "ci/cd",
    "continuous integration",
    "continuous deployment",

    # --------------------------------------------------------
    # CLOUD
    # --------------------------------------------------------

    "aws",
    "amazon web services",
    "azure",
    "microsoft azure",
    "gcp",
    "google cloud",

    # --------------------------------------------------------
    # DATA / AI / ML
    # --------------------------------------------------------

    "machine learning",
    "deep learning",
    "artificial intelligence",
    "tensorflow",
    "pytorch",
    "scikit-learn",
    "numpy",
    "pandas",
    "data analysis",
    "data science",
    "natural language processing",
    "nlp",

    # --------------------------------------------------------
    # PROGRAMMING CONCEPTS
    # --------------------------------------------------------

    "object oriented programming",
    "object-oriented programming",
    "oop",

    "data structures",
    "data structure",

    "algorithms",
    "algorithm",

    "dsa",

    "problem solving",
    "problem-solving",

    "debugging",
    "testing",

    "software development",
    "software engineering",

    # --------------------------------------------------------
    # OPERATING SYSTEM / NETWORKING
    # --------------------------------------------------------

    "linux",
    "unix",
    "windows",
    "networking",
    "computer networks",

    # --------------------------------------------------------
    # TOOLS
    # --------------------------------------------------------

    "visual studio",
    "visual studio code",
    "vs code",
    "jupyter",
    "postman",

    # --------------------------------------------------------
    # BUSINESS / ANALYTICS
    # --------------------------------------------------------

    "excel",
    "power bi",
    "tableau",

    # --------------------------------------------------------
    # SOFT SKILLS
    # --------------------------------------------------------

    "communication",
    "leadership",
    "teamwork",
    "time management",
    "analytical skills",
    "analytical thinking"
]


# ============================================================
# SKILL ALIASES
# ============================================================

SKILL_ALIASES = {

    "rest apis":
        "rest api",

    "restful api":
        "rest api",

    "restful apis":
        "rest api",

    "rest-api":
        "rest api",

    "restful-api":
        "rest api",

    "object-oriented programming":
        "object oriented programming",

    "problem-solving":
        "problem solving",

    "data structure":
        "data structures",

    "algorithm":
        "algorithms",

    "react.js":
        "react",

    "vue.js":
        "vue",

    "node.js":
        "node",

    "express.js":
        "express",

    "postgres":
        "postgresql",

    "amazon web services":
        "aws",

    "microsoft azure":
        "azure",

    "google cloud":
        "gcp",

    "scikit learn":
        "scikit-learn"
}


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text):

    if not text:
        return ""

    text = text.lower()

    text = text.replace("–", "-")
    text = text.replace("—", "-")
    text = text.replace("•", " ")
    text = text.replace("●", " ")

    # Normalize REST variations
    text = text.replace(
        "restful apis",
        "rest api"
    )

    text = text.replace(
        "restful api",
        "rest api"
    )

    text = text.replace(
        "rest apis",
        "rest api"
    )

    text = text.replace(
        "rest-api",
        "rest api"
    )

    text = text.replace(
        "restful-api",
        "rest api"
    )

    # Normalize OOP
    text = text.replace(
        "object-oriented programming",
        "object oriented programming"
    )

    # Normalize problem solving
    text = text.replace(
        "problem-solving",
        "problem solving"
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# NORMALIZE SKILL
# ============================================================

def normalize_skill(skill):

    skill = normalize_text(skill)

    if skill in SKILL_ALIASES:

        return SKILL_ALIASES[skill]

    return skill


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_path):

    if pdfplumber is None:

        return ""

    text = ""

    try:

        with pdfplumber.open(
            file_path
        ) as pdf:

            for page in pdf.pages:

                page_text = page.extract_text()

                if page_text:

                    text += "\n" + page_text

    except Exception as e:

        print(
            f"   PDF extraction error: {e}"
        )

    return text


# ============================================================
# PDF OCR
# ============================================================

def extract_pdf_ocr(file_path):

    if (
        convert_from_path is None
        or pytesseract is None
    ):

        return ""

    text = ""

    try:

        pages = convert_from_path(
            str(file_path),
            dpi=250
        )

        for page in pages:

            try:

                page_text = pytesseract.image_to_string(
                    page,
                    config="--psm 6"
                )

                if page_text:

                    text += "\n" + page_text

            except Exception:

                pass

    except Exception as e:

        print(
            f"   OCR error: {e}"
        )

    return text


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(file_path):

    if Document is None:

        return ""

    text = ""

    try:

        document = Document(
            file_path
        )

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text += (
                    "\n"
                    + paragraph.text
                )

        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():

                        text += (
                            "\n"
                            + cell.text
                        )

    except Exception as e:

        print(
            f"   DOCX extraction error: {e}"
        )

    return text


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_txt_text(file_path):

    try:

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            return file.read()

    except Exception as e:

        print(
            f"   TXT extraction error: {e}"
        )

        return ""


# ============================================================
# IMAGE OCR
# ============================================================

def extract_image_text(file_path):

    if (
        Image is None
        or pytesseract is None
    ):

        return ""

    try:

        image = Image.open(
            file_path
        )

        text = pytesseract.image_to_string(
            image,
            config="--psm 6"
        )

        return text

    except Exception as e:

        print(
            f"   Image OCR error: {e}"
        )

        return ""


# ============================================================
# GENERAL TEXT EXTRACTION
# ============================================================

def extract_text(file_path):

    extension = file_path.suffix.lower()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        print(
            "   Reading PDF..."
        )

        text = extract_pdf_text(
            file_path
        )

        if len(
            normalize_text(text)
        ) >= 100:

            print(
                "   Extraction method : PDF text"
            )

            return (
                text,
                "pdf_text"
            )

        print(
            "   Direct PDF text insufficient"
        )

        print(
            "   Extraction method : OCR"
        )

        text = extract_pdf_ocr(
            file_path
        )

        return (
            text,
            "pdf_ocr"
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif extension == ".docx":

        print(
            "   Reading DOCX..."
        )

        return (
            extract_docx_text(
                file_path
            ),
            "docx"
        )

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    elif extension == ".txt":

        print(
            "   Reading TXT..."
        )

        return (
            extract_txt_text(
                file_path
            ),
            "txt"
        )

    # --------------------------------------------------------
    # IMAGE
    # --------------------------------------------------------

    elif extension in {
        ".png",
        ".jpg",
        ".jpeg"
    }:

        print(
            "   Reading image with OCR..."
        )

        return (
            extract_image_text(
                file_path
            ),
            "ocr_image"
        )

    return (
        "",
        "unsupported"
    )


# ============================================================
# EMAIL EXTRACTION
# ============================================================

def extract_email(text):

    pattern = (
        r"[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\."
        r"[A-Za-z]{2,}"
    )

    matches = re.findall(
        pattern,
        text
    )

    if matches:

        return matches[0]

    return "Not detected"


# ============================================================
# PHONE EXTRACTION
# ============================================================

def extract_phone(text):

    patterns = [

        r"\+?\d[\d\s().-]{8,}\d",

        r"\b\d{10}\b"
    ]

    for pattern in patterns:

        matches = re.findall(
            pattern,
            text
        )

        for match in matches:

            digits = re.sub(
                r"\D",
                "",
                match
            )

            if 10 <= len(digits) <= 13:

                return match.strip()

    return "Not detected"


# ============================================================
# NAME EXTRACTION
# ============================================================

def extract_name(
    text,
    file_name
):

    lines = []

    for line in text.splitlines():

        line = line.strip()

        if not line:

            continue

        line = re.sub(
            r"\s+",
            " ",
            line
        )

        lines.append(line)

    ignored = {

        "resume",
        "curriculum vitae",
        "cv",
        "profile",
        "resume profile",
        "personal details",
        "contact",
        "career objective",
        "objective",
        "summary"
    }

    headings = {

        "education",
        "experience",
        "skills",
        "projects",
        "certifications",
        "achievements",
        "languages",
        "interests",
        "references",
        "professional experience",
        "work experience"
    }

    # --------------------------------------------------------
    # Look at first 20 lines
    # --------------------------------------------------------

    for line in lines[:20]:

        lower = line.lower().strip()

        if lower in ignored:

            continue

        if lower in headings:

            continue

        if "@" in line:

            continue

        if re.search(
            r"\d",
            line
        ):

            continue

        if len(line) < 3:

            continue

        if len(line) > 60:

            continue

        words = line.split()

        if not (
            2 <= len(words) <= 5
        ):

            continue

        valid = True

        for word in words:

            cleaned = re.sub(
                r"[^A-Za-z.'-]",
                "",
                word
            )

            if not cleaned:

                valid = False

                break

        if not valid:

            continue

        # Avoid obvious headings
        if any(
            heading in lower
            for heading in [
                "resume",
                "curriculum vitae",
                "developer",
                "engineer",
                "student",
                "profile",
                "objective"
            ]
        ) and len(words) <= 2:

            continue

        return line

    # --------------------------------------------------------
    # Filename fallback
    # --------------------------------------------------------

    name = Path(
        file_name
    ).stem

    name = re.sub(
        r"[_\-]+",
        " ",
        name
    )

    name = re.sub(
        r"\d+",
        " ",
        name
    )

    name = re.sub(
        r"\s+",
        " ",
        name
    ).strip()

    if 2 <= len(
        name.split()
    ) <= 5:

        return name.title()

    return "Not detected"


# ============================================================
# SKILL EXTRACTION
# ============================================================

def extract_skills(text):

    normalized = normalize_text(
        text
    )

    found = set()

    # --------------------------------------------------------
    # Check longer skills first
    # --------------------------------------------------------

    sorted_skills = sorted(
        SKILLS,
        key=len,
        reverse=True
    )

    for skill in sorted_skills:

        skill_normalized = normalize_skill(
            skill
        )

        # ----------------------------------------------------
        # Ignore generic API/APIs
        # ----------------------------------------------------

        if skill_normalized in {
            "api",
            "apis"
        }:

            continue

        # ----------------------------------------------------
        # Special handling for REST API
        # ----------------------------------------------------

        if skill_normalized == "rest api":

            rest_patterns = [

                r"\brest\s+api\b",

                r"\brestful\s+api\b",

                r"\brest\s+apis\b",

                r"\brestful\s+apis\b",

                r"\brest-api\b",

                r"\brestful-api\b"
            ]

            if any(
                re.search(
                    pattern,
                    text,
                    re.IGNORECASE
                )
                for pattern in rest_patterns
            ):

                found.add(
                    "rest api"
                )

            continue

        # ----------------------------------------------------
        # Object oriented programming
        # ----------------------------------------------------

        if skill_normalized == (
            "object oriented programming"
        ):

            patterns = [

                r"\bobject[\s-]+oriented[\s-]+programming\b",

                r"\boop\b"
            ]

            if any(
                re.search(
                    pattern,
                    text,
                    re.IGNORECASE
                )
                for pattern in patterns
            ):

                found.add(
                    "object oriented programming"
                )

            continue

        # ----------------------------------------------------
        # Problem solving
        # ----------------------------------------------------

        if skill_normalized == "problem solving":

            patterns = [

                r"\bproblem[\s-]+solving\b",

                r"\bproblem[\s-]+solving[\s-]+skills?\b"
            ]

            if any(
                re.search(
                    pattern,
                    text,
                    re.IGNORECASE
                )
                for pattern in patterns
            ):

                found.add(
                    "problem solving"
                )

            continue

        # ----------------------------------------------------
        # Normal matching
        # ----------------------------------------------------

        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(skill_normalized)
            + r"(?![a-z0-9])"
        )

        if re.search(
            pattern,
            normalized
        ):

            found.add(
                skill_normalized
            )

    return sorted(
        found
    )


# ============================================================
# EDUCATION EXTRACTION
# ============================================================

def extract_education(text):

    normalized = normalize_text(
        text
    )

    detected = []

    education_patterns = {

        "B.Tech": [

            r"\bb\.?\s*tech\b",

            r"\bbtech\b",

            r"\bbachelor\s+of\s+technology\b"
        ],

        "B.E": [

            r"\bb\.?\s*e\.?\b",

            r"\bbachelor\s+of\s+engineering\b"
        ],

        "M.Tech": [

            r"\bm\.?\s*tech\b",

            r"\bmtech\b",

            r"\bmaster\s+of\s+technology\b"
        ],

        "MCA": [

            r"\bmca\b",

            r"\bmaster\s+of\s+computer\s+applications\b"
        ],

        "BCA": [

            r"\bbca\b",

            r"\bbachelor\s+of\s+computer\s+applications\b"
        ],

        "B.Sc": [

            r"\bb\.?\s*sc\.?\b",

            r"\bbsc\b",

            r"\bbachelor\s+of\s+science\b"
        ],

        "M.Sc": [

            r"\bm\.?\s*sc\.?\b",

            r"\bmsc\b",

            r"\bmaster\s+of\s+science\b"
        ]

    }

    for degree, patterns in education_patterns.items():

        for pattern in patterns:

            if re.search(
                pattern,
                normalized
            ):

                if degree not in detected:

                    detected.append(
                        degree
                    )

                break

    # --------------------------------------------------------
    # Specializations
    # --------------------------------------------------------

    if (
        "computer science" in normalized
    ):

        detected.append(
            "Computer Science"
        )

    elif (
        "information technology"
        in normalized
    ):

        detected.append(
            "Information Technology"
        )

    elif (
        "electronics and communication"
        in normalized
    ):

        detected.append(
            "Electronics and Communication"
        )

    elif (
        "electronics and telecommunication"
        in normalized
    ):

        detected.append(
            "Electronics and Telecommunication"
        )

    if detected:

        return ", ".join(
            dict.fromkeys(
                detected
            )
        )

    return "Not detected"


# ============================================================
# EXPERIENCE EXTRACTION
# ============================================================

def extract_experience(text):

    normalized = normalize_text(
        text
    )

    values = []

    patterns = [

        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+of\s+experience",

        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+experience",

        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)\s*\+?\s*years?",

        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+in"
    ]

    for pattern in patterns:

        matches = re.findall(
            pattern,
            normalized
        )

        for match in matches:

            try:

                value = float(
                    match
                )

                if (
                    0 <= value <= 50
                ):

                    values.append(
                        value
                    )

            except Exception:

                pass

    if values:

        return max(
            values
        )

    # --------------------------------------------------------
    # Date based experience
    # --------------------------------------------------------

    years = re.findall(
        r"\b(19\d{2}|20\d{2})\b",
        normalized
    )

    year_values = []

    for year in years:

        year_int = int(
            year
        )

        if (
            1950 <= year_int <= 2035
        ):

            year_values.append(
                year_int
            )

    if len(
        year_values
    ) >= 2:

        earliest = min(
            year_values
        )

        latest = max(
            year_values
        )

        difference = (
            latest
            -
            earliest
        )

        if (
            0 <= difference <= 50
        ):

            return float(
                difference
            )

    return 0.0


# ============================================================
# PROJECT EXTRACTION
# ============================================================

def extract_projects(text):

    lines = text.splitlines()

    projects = []

    in_project_section = False

    stop_sections = {

        "education",
        "experience",
        "skills",
        "certifications",
        "achievements",
        "languages",
        "interests",
        "references",
        "work experience"
    }

    for line in lines:

        clean = line.strip()

        if not clean:

            continue

        lower = clean.lower()

        if (
            "project" in lower
            and len(clean) < 60
        ):

            in_project_section = True

            continue

        if in_project_section:

            if lower in stop_sections:

                break

            if len(clean) >= 10:

                projects.append(
                    clean
                )

            if len(projects) >= 10:

                break

    return projects


# ============================================================
# CANDIDATE DATA
# ============================================================

def extract_candidate_data(
    text,
    file_name
):

    return {

        "name":
            extract_name(
                text,
                file_name
            ),

        "email":
            extract_email(
                text
            ),

        "phone":
            extract_phone(
                text
            ),

        "skills":
            extract_skills(
                text
            ),

        "education":
            extract_education(
                text
            ),

        "experience_years":
            extract_experience(
                text
            ),

        "projects":
            extract_projects(
                text
            )
    }


# ============================================================
# SEMANTIC SIMILARITY
# ============================================================

def calculate_semantic_similarity(
    job_role,
    job_description,
    resume_text
):

    job_text = (
        job_role
        + " "
        + job_description
    )

    job_text = normalize_text(
        job_text
    )

    resume_text = normalize_text(
        resume_text
    )

    if (
        not job_text
        or not resume_text
    ):

        return 0.0

    try:

        vectorizer = TfidfVectorizer(

            stop_words="english",

            ngram_range=(1, 2),

            max_features=15000
        )

        matrix = vectorizer.fit_transform(
            [
                job_text,
                resume_text
            ]
        )

        similarity = cosine_similarity(
            matrix[0:1],
            matrix[1:2]
        )[0][0]

        return float(
            max(
                0,
                min(
                    1,
                    similarity
                )
            )
        )

    except Exception:

        return 0.0


# ============================================================
# SKILL MATCH
# ============================================================

def calculate_skill_match(
    required_skills,
    candidate_skills
):

    if not required_skills:

        return 1.0

    required = set(
        normalize_skill(
            skill
        )
        for skill in required_skills
    )

    candidate = set(
        normalize_skill(
            skill
        )
        for skill in candidate_skills
    )

    matched = (
        required &
        candidate
    )

    return (
        len(matched)
        /
        len(required)
    )


# ============================================================
# EXPERIENCE MATCH
# ============================================================

def calculate_experience_match(
    candidate_experience,
    required_experience
):

    if required_experience <= 0:

        return 1.0

    if (
        candidate_experience
        >=
        required_experience
    ):

        return 1.0

    if candidate_experience <= 0:

        return 0.0

    return (
        candidate_experience
        /
        required_experience
    )


# ============================================================
# EDUCATION MATCH
# ============================================================

def calculate_education_match(
    candidate_education,
    required_education
):

    if not required_education.strip():

        return 1.0

    candidate = normalize_text(
        candidate_education
    )

    required = normalize_text(
        required_education
    )

    if candidate == "not detected":

        return 0.0

    # --------------------------------------------------------
    # Required education can be:
    # B.Tech
    # B.Tech,B.E
    # B.Tech/B.E
    # B.Tech or B.E
    # --------------------------------------------------------

    requirements = re.split(
        r"[,;/|]+|\s+or\s+",
        required
    )

    requirements = [

        normalize_text(
            item
        )

        for item in requirements

        if normalize_text(item)
    ]

    for requirement in requirements:

        # B.Tech
        if (
            "b.tech" in requirement
            or "btech" in requirement
        ):

            if (
                "b.tech" in candidate
                or "b.e" in candidate
            ):

                return 1.0

        # B.E
        if (
            requirement == "b.e"
            or requirement == "be"
            or "b.e" in requirement
        ):

            if (
                "b.tech" in candidate
                or "b.e" in candidate
            ):

                return 1.0

        # MCA
        if "mca" in requirement:

            if "mca" in candidate:

                return 1.0

        # BCA
        if "bca" in requirement:

            if "bca" in candidate:

                return 1.0

        # M.Tech
        if (
            "m.tech" in requirement
            or "mtech" in requirement
        ):

            if "m.tech" in candidate:

                return 1.0

        # M.Sc
        if (
            "m.sc" in requirement
            or "msc" in requirement
        ):

            if "m.sc" in candidate:

                return 1.0

        # Direct comparison
        if requirement in candidate:

            return 1.0

    return 0.0


# ============================================================
# MATCH CANDIDATE
# ============================================================

def match_candidate(

    candidate,

    resume_text,

    job_role,

    job_description,

    required_skills,

    required_experience,

    required_education
):

    semantic = calculate_semantic_similarity(

        job_role,

        job_description,

        resume_text
    )

    skill_match = calculate_skill_match(

        required_skills,

        candidate["skills"]
    )

    experience_match = calculate_experience_match(

        candidate["experience_years"],

        required_experience
    )

    education_match = calculate_education_match(

        candidate["education"],

        required_education
    )

    # ========================================================
    # SCORE WEIGHTS
    # ========================================================

    semantic_weight = 0.40

    skill_weight = 0.30

    experience_weight = 0.15

    education_weight = 0.15

    overall = (

        semantic
        * semantic_weight

        +

        skill_match
        * skill_weight

        +

        experience_match
        * experience_weight

        +

        education_match
        * education_weight

    )

    # ========================================================
    # MATCHED / MISSING SKILLS
    # ========================================================

    required_set = set(

        normalize_skill(
            skill
        )

        for skill in required_skills
    )

    candidate_set = set(

        normalize_skill(
            skill
        )

        for skill in candidate["skills"]
    )

    matched_skills = sorted(
        required_set &
        candidate_set
    )

    missing_skills = sorted(
        required_set -
        candidate_set
    )

    # ========================================================
    # WHY CANDIDATE RECEIVED THIS SCORE
    # ========================================================

    reasons = []

    # --------------------------------------------------------
    # Semantic
    # --------------------------------------------------------

    if semantic >= 0.70:

        reasons.append(
            "Very strong semantic similarity with the job description"
        )

    elif semantic >= 0.50:

        reasons.append(
            "Strong semantic similarity with the job description"
        )

    elif semantic >= 0.30:

        reasons.append(
            "Moderate semantic similarity with the job description"
        )

    else:

        reasons.append(
            "Low semantic similarity with the job description"
        )

    # --------------------------------------------------------
    # Skills
    # --------------------------------------------------------

    if required_skills:

        reasons.append(

            f"Matches {len(matched_skills)} "
            f"of {len(required_skills)} required skills"
        )

    else:

        reasons.append(
            "No specific required skills were provided"
        )

    # --------------------------------------------------------
    # Experience
    # --------------------------------------------------------

    if required_experience > 0:

        if (
            candidate["experience_years"]
            >=
            required_experience
        ):

            reasons.append(
                "Meets the required experience"
            )

        else:

            reasons.append(
                "Does not fully meet the required experience"
            )

    else:

        reasons.append(
            "No specific experience requirement was provided"
        )

    # --------------------------------------------------------
    # Education
    # --------------------------------------------------------

    if required_education:

        if education_match >= 1:

            reasons.append(
                "Meets the required education"
            )

        else:

            reasons.append(
                "Required education was not detected"
            )

    else:

        reasons.append(
            "No specific education requirement was provided"
        )

    # ========================================================
    # RECOMMENDATION
    # ========================================================

    if overall >= 0.75:

        recommendation = (
            "HIGH SUITABILITY"
        )

    elif overall >= 0.55:

        recommendation = (
            "GOOD SUITABILITY"
        )

    elif overall >= 0.35:

        recommendation = (
            "MODERATE SUITABILITY"
        )

    else:

        recommendation = (
            "LOW SUITABILITY"
        )

    return {

        "overall_score":
            round(
                overall * 100,
                2
            ),

        "semantic_similarity":
            round(
                semantic * 100,
                2
            ),

        "skill_match":
            round(
                skill_match * 100,
                2
            ),

        "experience_match":
            round(
                experience_match * 100,
                2
            ),

        "education_match":
            round(
                education_match * 100,
                2
            ),

        "matched_skills":
            matched_skills,

        "missing_skills":
            missing_skills,

        "reasons":
            reasons,

        "recommendation":
            recommendation
    }


# ============================================================
# JOB INPUT
# ============================================================

def get_job_details():

    print()
    print("=" * 70)
    print("JOB DETAILS")
    print("=" * 70)

    job_role = input(
        "Job Role: "
    ).strip()

    print()
    print(
        "Enter Job Description."
    )

    print(
        "Paste the complete JD."
    )

    print(
        "Type END on a NEW LINE when finished."
    )

    print()

    jd_lines = []

    while True:

        line = input()

        if line.strip().upper() == "END":

            break

        jd_lines.append(
            line
        )

    job_description = "\n".join(
        jd_lines
    ).strip()

    # --------------------------------------------------------
    # Experience
    # --------------------------------------------------------

    while True:

        experience_input = input(

            "\nRequired Experience in years "
            "(enter 0 if not required): "

        ).strip()

        try:

            required_experience = float(
                experience_input
            )

            if required_experience < 0:

                print(
                    "Enter 0 or a positive number."
                )

                continue

            break

        except ValueError:

            print(
                "Please enter a valid number."
            )

    # --------------------------------------------------------
    # Education
    # --------------------------------------------------------

    required_education = input(

        "Required Education "
        "(example: B.Tech,B.E,MCA or ENTER for none): "

    ).strip()

    return (

        job_role,

        job_description,

        required_experience,

        required_education
    )


# ============================================================
# DISPLAY CANDIDATE
# ============================================================

def display_candidate(
    rank,
    result
):

    print()
    print("=" * 70)

    print(
        f"RANK #{rank}"
    )

    print("=" * 70)

    print(
        f"Candidate Name : "
        f"{result['candidate']['name']}"
    )

    print(
        f"Resume         : "
        f"{result['resume_file']}"
    )

    print(
        f"Email          : "
        f"{result['candidate']['email']}"
    )

    print(
        f"Phone          : "
        f"{result['candidate']['phone']}"
    )

    print()

    # --------------------------------------------------------
    # SCORE
    # --------------------------------------------------------

    print(
        f"Overall Score       : "
        f"{result['overall_score']:.2f}%"
    )

    print(
        f"Semantic Similarity : "
        f"{result['semantic_similarity']:.2f}%"
    )

    print(
        f"Skill Match         : "
        f"{result['skill_match']:.2f}%"
    )

    print(
        f"Experience Match    : "
        f"{result['experience_match']:.2f}%"
    )

    print(
        f"Education Match     : "
        f"{result['education_match']:.2f}%"
    )

    print()

    # --------------------------------------------------------
    # EXPERIENCE / EDUCATION
    # --------------------------------------------------------

    print(
        f"Experience : "
        f"{result['candidate']['experience_years']:.1f} years"
    )

    print(
        f"Education  : "
        f"{result['candidate']['education']}"
    )

    print()

    # --------------------------------------------------------
    # MATCHED SKILLS
    # --------------------------------------------------------

    print(
        "Matched Skills:"
    )

    if result["matched_skills"]:

        for skill in result["matched_skills"]:

            print(
                f"  ✓ {skill}"
            )

    else:

        print(
            "  None"
        )

    print()

    # --------------------------------------------------------
    # MISSING SKILLS
    # --------------------------------------------------------

    print(
        "Missing Skills:"
    )

    if result["missing_skills"]:

        for skill in result["missing_skills"]:

            print(
                f"  - {skill}"
            )

    else:

        print(
            "  None"
        )

    print()

    # --------------------------------------------------------
    # REASON
    # --------------------------------------------------------

    print(
        "WHY THIS CANDIDATE WAS RANKED HERE:"
    )

    for reason in result["reasons"]:

        print(
            f"  ✓ {reason}"
        )

    print()

    print(
        f"RECOMMENDATION: "
        f"{result['recommendation']}"
    )

    print("=" * 70)


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 70)
    print("HR RESUME AI - CANDIDATE MATCHING")
    print("=" * 70)

    # --------------------------------------------------------
    # Check resume folder
    # --------------------------------------------------------

    if not RESUME_FOLDER.exists():

        print()

        print(
            "ERROR: Resume folder does not exist:"
        )

        print(
            RESUME_FOLDER
        )

        return

    # --------------------------------------------------------
    # Get job details
    # --------------------------------------------------------

    (

        job_role,

        job_description,

        required_experience,

        required_education

    ) = get_job_details()

    if not job_role:

        print(
            "\nERROR: Job role cannot be empty."
        )

        return

    if not job_description:

        print(
            "\nERROR: Job description cannot be empty."
        )

        return

    # --------------------------------------------------------
    # Extract required skills
    # --------------------------------------------------------

    required_skills = extract_skills(
        job_description
    )

    # --------------------------------------------------------
    # Display requirements
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("JOB REQUIREMENTS")
    print("=" * 70)

    print(
        f"Job Role: {job_role}"
    )

    print()

    print(
        "Job Description:"
    )

    print(
        job_description
    )

    print()

    print(
        "Required Experience:"
    )

    if required_experience > 0:

        print(
            f"  {required_experience} years"
        )

    else:

        print(
            "  No experience requirement"
        )

    print()

    print(
        "Required Education:"
    )

    if required_education:

        print(
            f"  {required_education}"
        )

    else:

        print(
            "  None"
        )

    print()

    print(
        "Required Skills:"
    )

    if required_skills:

        for skill in required_skills:

            print(
                f"  • {skill}"
            )

    else:

        print(
            "  None detected"
        )

    # --------------------------------------------------------
    # Find resumes
    # --------------------------------------------------------

    resume_files = []

    for file_path in RESUME_FOLDER.iterdir():

        if not file_path.is_file():

            continue

        if (
            file_path.suffix.lower()
            not in SUPPORTED_EXTENSIONS
        ):

            continue

        resume_files.append(
            file_path
        )

    resume_files.sort(
        key=lambda x: x.name.lower()
    )

    print()
    print("=" * 70)

    print(
        f"RESUMES FOUND: "
        f"{len(resume_files)}"
    )

    print("=" * 70)

    if not resume_files:

        print(
            "No resumes found."
        )

        return

    results = []

    # ========================================================
    # PROCESS RESUMES
    # ========================================================

    for index, file_path in enumerate(

        resume_files,

        start=1
    ):

        print()

        print(
            f"[{index}/{len(resume_files)}] "
            f"{file_path.name}"
        )

        print(
            "-" * 70
        )

        try:

            text, extraction_method = extract_text(
                file_path
            )

            normalized = normalize_text(
                text
            )

            characters = len(
                normalized
            )

            words = len(
                normalized.split()
            )

            print(
                f"   Characters : {characters}"
            )

            print(
                f"   Words      : {words}"
            )

            if characters < 30:

                print(
                    "   ⚠ Insufficient extracted text"
                )

                print(
                    "   Candidate skipped"
                )

                continue

            # ------------------------------------------------
            # Extract candidate information
            # ------------------------------------------------

            candidate = extract_candidate_data(

                text,

                file_path.name
            )

            # ------------------------------------------------
            # Match candidate
            # ------------------------------------------------

            match = match_candidate(

                candidate,

                text,

                job_role,

                job_description,

                required_skills,

                required_experience,

                required_education
            )

            result = {

                "resume_file":
                    file_path.name,

                "candidate":
                    candidate,

                "overall_score":
                    match["overall_score"],

                "semantic_similarity":
                    match["semantic_similarity"],

                "skill_match":
                    match["skill_match"],

                "experience_match":
                    match["experience_match"],

                "education_match":
                    match["education_match"],

                "matched_skills":
                    match["matched_skills"],

                "missing_skills":
                    match["missing_skills"],

                "reasons":
                    match["reasons"],

                "recommendation":
                    match["recommendation"],

                "extraction_method":
                    extraction_method,

                "text_characters":
                    characters,

                "text_words":
                    words
            }

            results.append(
                result
            )

            print(
                f"   Candidate : "
                f"{candidate['name']}"
            )

            print(
                f"   Score     : "
                f"{match['overall_score']:.2f}%"
            )

        except Exception as e:

            print(
                f"   ✗ Error: {e}"
            )

    # ========================================================
    # SORT CANDIDATES
    # ========================================================

    results.sort(

        key=lambda item: (

            item["overall_score"],

            item["skill_match"],

            item["semantic_similarity"],

            item["experience_match"],

            item["education_match"]
        ),

        reverse=True
    )

    # ========================================================
    # ASSIGN RANK
    # ========================================================

    for rank, result in enumerate(

        results,

        start=1
    ):

        result["rank"] = rank

    # ========================================================
    # FINAL RANKING
    # ========================================================

    print()
    print()
    print("=" * 70)
    print("FINAL CANDIDATE RANKING")
    print("=" * 70)

    for result in results:

        display_candidate(

            result["rank"],

            result
        )

    # ========================================================
    # SAVE RESULTS
    # ========================================================

    output = {

        "job_details": {

            "job_role":
                job_role,

            "job_description":
                job_description,

            "required_experience":
                required_experience,

            "required_education":
                required_education,

            "required_skills":
                required_skills
        },

        "total_candidates_analyzed":
            len(results),

        "candidates":
            results
    }

    try:

        with open(

            RESULT_FILE,

            "w",

            encoding="utf-8"

        ) as file:

            json.dump(

                output,

                file,

                indent=4,

                ensure_ascii=False
            )

    except Exception as e:

        print(
            f"Could not save results: {e}"
        )

        return

    # ========================================================
    # FINISHED
    # ========================================================

    print()

    print("=" * 70)

    print(
        "MATCHING COMPLETED"
    )

    print("=" * 70)

    print(
        f"Candidates analyzed: "
        f"{len(results)}"
    )

    print()

    print(
        "Results saved to:"
    )

    print(
        RESULT_FILE
    )

    print("=" * 70)


# ============================================================
# START PROGRAM
# ============================================================

if __name__ == "__main__":

    main()